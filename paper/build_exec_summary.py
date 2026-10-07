"""One-page executive summary of the SPACOMM 2027 submission."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for s in doc.sections:
    s.top_margin = Pt(36); s.bottom_margin = Pt(36)
    s.left_margin = Pt(54); s.right_margin = Pt(54)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)

def para(text, bold=False, size=11, center=False, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if center:
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.name = "Times New Roman"
    return p

para("EXECUTIVE SUMMARY", bold=True, size=13, center=True, space_after=2)
para("From Illustrative to Measured: Task-Oriented Semantic mIoT over NTN "
     "for Remote Agriculture and Environmental Sensing", bold=True, center=True, space_after=2)
para("Rao Yenamandra (DigitalTwinSim, San Diego, CA, USA) \u00b7 "
     "Mubanga Nsofu (Vodacom, Johannesburg, South Africa) \u00b7 "
     "Ram Asokan (Wireless Research Center of North Carolina, USA)",
     size=10, center=True, space_after=2)
para("Target: SPACOMM 2027 (IARIA), Lisbon, April 18\u201322, 2027 \u00b7 "
     "Submission deadline: December 28, 2026 \u00b7 Status: manuscript complete (8 pp.), not submitted",
     size=10, center=True, space_after=8)

para("THE PROBLEM", bold=True, size=11, space_after=2)
para("An earlier IEEE submission of this work was rejected: reviewers found no substantive "
     "experiments \u2014 the headline 32.5\u00d7 gain was an illustrative calculation, a section titled "
     "\u201cRequired Results\u201d made the paper look unfinished, five applications were too broad, and "
     "template/proofreading defects hurt credibility. This revision answers every criticism with "
     "measured evidence.", space_after=6)

para("WHAT WAS DONE", bold=True, size=11, space_after=2)
para("The paper was rewritten from scratch around a three-stage, trace-driven simulation "
     "(seeded, reproducible, 5 held-out evaluation seeds) comparing semantic/task-oriented reporting "
     "against raw periodic, compressed periodic, and event-triggered baselines. Scope was narrowed "
     "to two India-oriented applications: (1) irrigation decisions for smart agriculture (soil-moisture "
     "traces, 6-hour deadline) and (2) flood warnings for environmental monitoring (river-level traces, "
     "45-minute deadline). Link and energy models are anchored in published NTN field measurements, "
     "and the architecture figure was completely redesigned.", space_after=6)

para("HEADLINE RESULTS (all measured, mean \u00b1 sd over 5 seeds)", bold=True, size=11, space_after=2)
bullets = [
    "Report suppression dominates: of the 42.6\u00d7 offered-byte gain, 18.9\u00d7 comes from sending "
    "fewer reports (only ~5% admitted) and just 2.25\u00d7 from smaller payloads. The 60 B header caps "
    "the airtime gain at 19.6\u00d7.",
    "Capacity: at 95% deadline reliability, semantic reporting supports 1,600 devices/beam for "
    "irrigation vs 50 for raw periodic. For flood warnings, raw and compressed reporting collapse "
    "entirely (0 supportable devices); semantic supports 400.",
    "Task quality: flood-warning recall 0.928 (semantic) vs 0.157 (raw, which misses 84% of floods) "
    "vs 0.657 (event-triggered). Task-oriented gain: 104.6\u00d7 (floods), 19.6\u00d7 (irrigation).",
    "Energy: ~300 J vs ~28 kJ per device per 90 days (\u224899% saving); the honest uncongested "
    "figure is ~95%, since part of raw's cost is a congestion retry-storm.",
    "Adaptation: a drift detector plus one threshold update restores flood-warning recall from "
    "0.746 (stale dry-season tuning) to 1.000.",
]
for b in bullets:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(b); r.font.size = Pt(11); r.font.name = "Times New Roman"

para("NEXT STEPS (framed as future work in the paper)", bold=True, size=11, space_after=2)
para("Validation on field sensor datasets (rainfall, soil-moisture, river-level records); "
     "multi-constellation link sweeps; NTN-specific random-access (NPRACH) modeling and hardware "
     "energy measurements on MCU-class devices; heterogeneous device/receiver decision rules; "
     "alternative deadline/reliability operating points. All reported gains are resource and "
     "task-outcome gains bounded to the evaluated applications \u2014 no Shannon-capacity claims.",
     space_after=6)

para("REMAINING STEPS", bold=True, size=11, space_after=2)
para("1) Confirm affiliations as listed above. 2) Decide: submit at 8 pages (2 pages over the "
     "included 6 \u2014 extra-page fees apply) or cut to 6. 3) Sanity-check the six flagged simulation "
     "assumptions. 4) Submit via the SPACOMM 2027 system before December 28, 2026.",
     space_after=0)

out = "/home/hatch/workspace/ntn-eval/paper/Executive_Summary_SPACOMM2027.docx"
doc.save(out)
print("saved", out)
