"""One-page cover mail to co-authors for the SPACOMM 2027 revision."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for s in doc.sections:
    s.top_margin = Pt(48); s.bottom_margin = Pt(48)
    s.left_margin = Pt(64); s.right_margin = Pt(64)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(11)

def para(text, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.name = "Times New Roman"
    return p

para("Subject: For your review \u2014 revised NTN semantic-communications paper (SPACOMM 2027)",
     bold=True, size=11, space_after=10)
para("Dear Mubanga and Ram,", space_after=6)
para("You will recall our paper \u201cApplication-Specific Semantic and Task-Oriented Communication "
     "for Massive IoT over NTN,\u201d which IEEE rejected. The reviewers\u2019 main objection was the "
     "absence of substantive experimental results \u2014 the headline 32.5\u00d7 gain was an illustrative "
     "calculation, a section titled \u201cRequired Results\u201d made the paper look unfinished, and five "
     "applications were too broad.", space_after=6)
para("I have revised it from scratch for resubmission to SPACOMM 2027 (IARIA, Lisbon, "
     "April 18\u201322, 2027; full-paper deadline December 28, 2026), a non-IEEE venue with a "
     "dedicated satellite\u2013IoT\u20136G track. The retitled paper is \u201cFrom Illustrative to Measured: "
     "Task-Oriented Semantic mIoT over NTN for Remote Agriculture and Environmental Sensing.\u201d",
     space_after=6)
para("What changed:", bold=True, space_after=3)
bullets = [
    "The forward-looking \u201cRequired Results\u201d section is gone, replaced by a real three-stage, "
    "trace-driven simulation (seeded and reproducible) comparing semantic reporting against raw "
    "periodic, compressed periodic, and event-triggered baselines.",
    "Scope narrowed to two India-oriented applications: irrigation decisions (smart agriculture) "
    "and flood warnings (environmental monitoring).",
    "Measured headline results: 42.6\u00d7 offered-byte gain \u2014 18.9\u00d7 from report suppression, "
    "only 2.25\u00d7 from smaller payloads; supportable devices per beam 50 \u2192 1,600 for irrigation; "
    "conventional periodic reporting cannot meet the 45-minute flood-warning deadline at any modeled "
    "population (semantic supports 400); flood-warning recall 0.928 vs 0.157 for raw; ~99% device "
    "energy saving.",
    "All figures regenerated (architecture figure completely redesigned); 26 references; "
    "reformatted to the IARIA template; limitations reframed as future work.",
]
for b in bullets:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(b); r.font.size = Pt(11); r.font.name = "Times New Roman"

para("What I would like from you: please read the attached manuscript and send your edits and "
     "comments. In particular, please confirm your affiliation line as listed below, and flag "
     "anything I have missed or misstated.", space_after=6)
para("Affiliations as currently listed: Rao Yenamandra \u2014 DigitalTwinSim, San Diego, CA, USA; "
     "Mubanga Nsofu \u2014 Vodacom, Johannesburg, South Africa; Ram Asokan \u2014 Wireless Research "
     "Center of North Carolina, USA.", space_after=6)
para("If you could get comments back by [DATE], that leaves comfortable time for revisions before "
     "the December 28 submission deadline. Attached: the manuscript (PDF) and a one-page executive "
     "summary.", space_after=6)
para("Warm regards,", space_after=2)
para("Rao", space_after=0)

out = "/home/hatch/workspace/ntn-eval/paper/Cover_Mail_Coauthors_SPACOMM2027.docx"
doc.save(out)
print("saved", out)
